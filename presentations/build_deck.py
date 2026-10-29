"""Author an editable Chinese technical deck from the paper and actual user CSV."""
from pathlib import Path
import csv,collections,json,math,os,shutil,re
from PIL import Image,ImageFont
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE,XL_LEGEND_POSITION,XL_MARKER_STYLE,XL_TICK_LABEL_POSITION
from pptx.chart.data import XyChartData
from pptx.enum.text import MSO_ANCHOR

R=Path(__file__).resolve().parent
A=R/'assets'
REPO=R.parent
SOURCE=A/'s11_SLOT_source.csv'
NAVY='#11283F';INK='#183248';TEAL='#087F8C';GOLD='#E8B545';MUTED='#567080';BG='#F7F9FB';LINE='#D7E1E8';WHITE='#FFFFFF';PINK='#A33C6D'
FONT='Noto Sans CJK SC'
FP='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
FB='/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc'
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
slides_meta=[];checks=[]

def rgb(s):return RGBColor.from_string(s.lstrip('#'))
def rect(s,x,y,w,h,fill=WHITE,line=None,rounded=False):
 o=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 o.fill.solid();o.fill.fore_color.rgb=rgb(fill)
 o.line.fill.background() if line is None else None
 if line:o.line.color.rgb=rgb(line)
 return o

def wrap(t,size,width,bold=False):
 font=ImageFont.truetype(FB if bold else FP,int(size*4),index=2)
 limit=width*72*4
 result=[]
 for para in t.split('\n'):
  tokens=re.findall(r'[A-Za-z0-9][A-Za-z0-9._/+-]*|[^A-Za-z0-9]',para)
  current=[]
  for token in tokens:
   candidate=''.join(current)+token
   if current and font.getlength(candidate)>limit:
    if token in '，。；：！？、）】' and len(current)>1:
     last=current.pop()
     result.append(''.join(current));current=[last,token]
    else:
     result.append(''.join(current).rstrip());current=[token.lstrip()]
   else:current.append(token)
  result.append(''.join(current))
 return result

def text(s,x,y,w,h,t,size=20,color=INK,bold=False):
 lines=wrap(str(t),size,w-.05,bold)
 needed=len(lines)*size*1.24/72
 if needed>h+.025:
  raise ValueError('Text overflow: '+str(t)[:50]+' needs '+str(round(needed,2))+' vs '+str(h))
 o=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
 tf=o.text_frame;tf.clear();tf.word_wrap=False
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(lines):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
  p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
  p.line_spacing=Pt(size*1.24);p.space_before=Pt(0);p.space_after=Pt(0)
 checks.append({'text':str(t),'x':x,'y':y,'w':w,'h':h,'size':size,'lines':len(lines)})
 return o

def new(title,section='仿真思路',subtitle='',source='依据：用户提供论文 arXiv:2603.03554v3；解释见备注。',notes=''):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
 rect(s,0,0,.15,7.5,TEAL)
 text(s,.48,.18,11,.28,section,11,TEAL,True)
 text(s,.48,.6,12.3,.55,title,28,INK,True)
 if subtitle:text(s,.5,1.23,12.2,.35,subtitle,14,MUTED)
 rect(s,.5,1.71,12.3,.014,LINE)
 text(s,.5,7.11,11.7,.2,source,8.5,MUTED)
 text(s,12.1,7.07,.65,.3,'%02d'%(len(prs.slides)),12,MUTED)
 s.notes_slide.notes_text_frame.text=(notes or '讲解时按页面逻辑展开，区分原文陈述、理论解释和当前模型假设。')+'\n\n'+source
 slides_meta.append({'page':len(prs.slides),'title':title,'section':section,'source':source,'notes':notes})
 return s

def callout(s,t,y=6.43,color=TEAL):
 rect(s,.55,y,12.15,.48,'#E5F2F3')
 text(s,.73,y+.07,11.8,.33,t,15,color,True)

def bullets(s,items,x=.72,y=2.0,w=11.8,gap=.95,size=21):
 for i,item in enumerate(items):
  text(s,x,y+i*gap,.3,.4,'•',size,TEAL,True)
  text(s,x+.34,y+i*gap,w-.34,gap-.12,item,size)

def cards(title,items,section='结构原理',subtitle='',source=None,notes=''):
 s=new(title,section,subtitle,**({'source':source} if source else {}),notes=notes)
 n=len(items);w=(12.1-.25*(n-1))/n
 for i,(head,body) in enumerate(items):
  x=.6+i*(w+.25);rect(s,x,1.98,w,4.14,WHITE,LINE,True)
  text(s,x+.2,2.2,w-.4,.7,head,21,TEAL,True)
  text(s,x+.2,3.08,w-.4,2.8,body,19)
 return s

def pic(s,path,x,y,w,h):
 im=Image.open(path);a,b=im.size;k=min(w/a,h/b)
 return s.shapes.add_picture(str(path),Inches(x+(w-a*k)/2),Inches(y+(h-b*k)/2),width=Inches(a*k),height=Inches(b*k))

def image_slide(title,path,items,section='结构原理',source=None,notes=''):
 s=new(title,section,**({'source':source} if source else {}),notes=notes)
 rect(s,.55,1.95,8.05,4.25,WHITE,LINE)
 pic(s,path,.65,2.05,7.85,4.05)
 bullets(s,items,x=8.85,y=2.12,w=3.75,gap=1.29,size=18)
 return s

def table(s,headers,rows,x=.62,y=2.0,width=12.08,row_h=.56,col_widths=None,size=17):
 widths=col_widths or [width/len(headers)]*len(headers)
 xx=x
 for k,h in enumerate(headers):
  rect(s,xx,y,widths[k],row_h,INK);text(s,xx+.12,y+.1,widths[k]-.24,row_h-.12,h,size,WHITE,True);xx+=widths[k]
 for i,row in enumerate(rows):
  xx=x
  for j,v in enumerate(row):
   yy=y+(i+1)*row_h;rect(s,xx,yy,widths[j],row_h,WHITE if i%2==0 else '#EDF3F6')
   text(s,xx+.12,yy+.1,widths[j]-.24,row_h-.12,str(v),size)
   xx+=widths[j]

def flow(title,steps,bottom,section='仿真方法',notes=''):
 s=new(title,section,notes=notes)
 n=len(steps);w=(11.7-.35*(n-1))/n
 for i,(a,b) in enumerate(steps):
  x=.7+i*(w+.35)
  rect(s,x,2.4,w,2.8,WHITE,LINE,True)
  text(s,x+.16,2.62,w-.32,.75,a,22,TEAL,True)
  text(s,x+.16,3.65,w-.32,1.3,b,18)
  if i<n-1:text(s,x+w+.03,3.42,.3,.4,'→',22,TEAL,True)
 callout(s,bottom)
 return s

# Reviewed uploaded CSV.
g=collections.defaultdict(list)
with SOURCE.open(encoding='utf-8-sig',newline='') as f:
 rd=csv.reader(f);next(rd)
 for row in rd:
  l,freq,s=map(float,row);g[l].append((freq,s))
summary=[]
for l,rows in sorted(g.items()):
 rows.sort();assert len(rows)==241 and len(set(f for f,s in rows))==241
 assert [f for f,s in rows]==[100+.5*i for i in range(241)]
 minima=[(f,s) for i,(f,s) in enumerate(rows) if 0<i<len(rows)-1 and s<rows[i-1][1] and s<=rows[i+1][1]]
 high=min((v for v in minima if 140<=v[0]<=200),key=lambda x:x[1])
 low=min((v for v in minima if 125<=v[0]<=135),key=lambda x:x[1])
 summary.append((l,high[0],high[1],low[0],low[1],dict(rows)[130.]))
with (R/'slot_sweep_summary.csv').open('w',newline='',encoding='utf-8') as f:
 w=csv.writer(f);w.writerow(['slot_um','high_trough_GHz','high_S11_dB','near130_GHz','near130_S11_dB','S11_at130_dB']);w.writerows(summary)

# 01
s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(NAVY)
text(s,.65,.6,11,.4,'论文解读 × HFSS 2021 R1 × 仿真验证',15,GOLD,True)
text(s,.65,1.4,11.8,1.7,'石墨烯零偏压太赫兹探测器\n结构设计与仿真思路',37,WHITE,True)
text(s,.7,3.45,11.9,1.0,'从双槽天线、阻抗匹配与齿形接触，\n到可追溯的建模、扫参与论文数据核对',23,'#C7DCE4')
text(s,.7,5.25,11.8,.8,'基于用户提供论文与 s11 SLOT.csv\n包含已跑通基线、旧模型局限与新版结构核对流程',17,'#C7DCE4')
text(s,.7,6.8,11.8,.3,'范围声明：不是“论文数据已复现”的结题报告。',14,GOLD)
s.notes_slide.notes_text_frame.text='本讲义面向 HFSS 初学者。论文为 arXiv:2603.03554v3。原矩形地模型在用户本机建立并完成过求解；新版有限 GSG 版图只做了云端二维几何检查，尚未执行 HFSS。槽长扫描数据属于旧模型，不能挪作新版或论文的结果。'
slides_meta.append({'page':1,'title':'石墨烯零偏压太赫兹探测器：结构设计与仿真思路','section':'封面','source':'用户上传论文与CSV','notes':'严格区分已验证与未验证内容。'})
# 02
s=cards('先分清：论文事实、模型假设、用户结果',[
 ('原文明确给出','双槽与扼流拓扑；约1 kΩ石墨烯；300 nm SiO₂；Cr/Au厚度；图2d为计算结果。'),
 ('当前模型假设','多数尺寸来自图片；端口参考1000 Ω；PEC；开放端读出线；衬底与透镜仍有简化。'),
 ('已有计算证据','旧矩形地模型在用户HFSS中跑通；已有7组槽长CSV；新版有限GSG结构尚未求解。')],section='阅读约定',notes='不要把结构看起来相似、设计检查通过或一个反射低谷接近130 GHz等同于复现。后续所有曲线页均标注来自旧模型的用户CSV。')
callout(s,'本PPT的目标：解释“为何这样建模”和“如何判断结果可信”。')
# 03
flow('整项工作分三层，每层回答不同问题',[
 ('电磁层','天线接收、馈点阻抗、辐射与材料损耗'),('热电层','石墨烯吸收、电子温升与光热电压'),('高速读出层','RC、传输线、封装与差频输出')],
 '图2d主要属于电磁层；仅靠HFSS端口S11不能验证完整探测性能。',section='整体框架')
# 04
image_slide('把原图拆成五个功能区',A/'antenna_detail.png',[
 '左：G–S–G焊盘与共面读出线。','中：交替宽窄段，形成扼流网络。','右：双槽、局部接触与石墨烯敏感区。'],source='来源：论文图1a，原图裁切；不是本次仿真场图。',notes='图中的小插图是独立放大的场示意，不与大图使用同一尺寸比例。读取500 μm标尺只能约束大图宏观尺寸。')
# 05
s=new('130 GHz 与“超过43 GHz”不是同一带宽','频率概念',notes='正文用双源差频产生IF，IF等于两源频率差。43GHz受测量链路限制，不是入射天线的中心频率，也不能从100–220GHz的S11直接推出。')
table(s,['频率尺度','在文中的含义','对应的仿真/测量'],[
 ['约130 GHz','太赫兹载频、参考照射频率','天线、局部场和吸收'],
 ['100–220 GHz','图2d显示的载频扫描范围','Zin与S11的频谱'],
 ['IF = |f₁ − f₂|','双源混频产生的电学输出频率','光电响应与读出链路'],
 ['>43 GHz','论文测得的电学调制/差频带宽下限','整机高速测量，受仪器限制']],row_h=.82,col_widths=[2.5,4.3,5.28],size=18)
# 06
s=cards('为什么微小石墨烯需要天线？',[
 ('尺度不匹配','130 GHz自由空间波长约2.31 mm，而石墨烯敏感区为微米量级。直接接收自由空间辐射效率受限。'),
 ('把能量送到局部','金属开口与馈电结构产生局部增强，使辐射与敏感区域耦合。增强场必须出现在有效吸收区域。'),
 ('不能只看最大场值','尖角的数值峰值可能依赖网格。应比较区域积分吸收、实际馈点功率及收敛后的场分布。')],notes='波长由c/f计算。场增强不是自动等同于吸收率提升；还需材料复电导和实际体积/面积。')
# 07
flow('能量与信号沿着两条路径传播',[
 ('接收','自由空间THz\n→透镜/衬底'),('耦合','双槽与馈线\n→石墨烯附近'),('转换','电子加热\n→PTE电信号'),('读出','共面线/PCB\n→放大器与仪器')],
 '需要同时减少太赫兹能量泄漏和读出链路的寄生损失。',section='结构原理')
# 08
s=cards('双槽：收集辐射，也改变馈点阻抗',[
 ('槽内电场','金属开口两侧形成电势差，电场集中到槽与馈点附近；这是作者讨论的高电阻耦合机制。'),
 ('结构与电长度','正文描述半波长槽，以及探测器到槽的四分之一波长距离；实际有效波长取决于传播模式。'),
 ('仿真要验证','同时看场的位置、表面电流、馈点阻抗与功率。槽距和馈线的变化可能引起耦合模式变化。')],source='依据：论文第2–4页；文献[20]、[21]为本文引用的双槽天线依据。',notes='本文参考[20]为Filipovic等1993年MTT文章，[21]为Zhang等2011年TST文章。本PPT未独立核查这两篇全文，不把它们当作当前几何精确尺寸来源。')
# 09
s=new('匹配对象约为1 kΩ，而不是默认50 Ω','阻抗匹配',notes='用天线的Thevenin等效解释功率传输。在实数负载近似下，天线端口阻抗接近1000+j0Ω有利于向1kΩ负载传能。真实石墨烯若含电抗，应使用共轭匹配。')
text(s,.75,2.05,11.7,.7,'理想功率传输条件：Z天线 ≈ Z负载*',28,TEAL,True)
bullets(s,['论文指出石墨烯在电中性点附近约为1 kΩ；普通低阻抗天线与之连接会产生失配。','高阻抗天线是作者选择的方向；不是先把石墨烯强行降到50 Ω。','必须记录端口参考阻抗。以50 Ω和1000 Ω归一化得到的S11，不能直接比较。'],y=3.2,gap=.88,size=21)
callout(s,'参考阻抗改变会改变S11；物理结构与终端条件不变时，正确换算的Zin应一致。')
# 10
s=new('λ/4 变换说明设计方向，但不能代替全波模型','传输线依据',notes='公式适用于均匀无损线、指定负载、给定传播常数。天线槽、渐变、有限地与多段扼流不等于一段理想传输线。不能用自由空间λ直接确定实际尺寸，也不能把图片估计62.5μm的距离认定为已实现作者λ/4条件。')
text(s,.75,2.04,11.7,.6,'在 βl = π/2 时：Zin = Zt² / ZL',29,TEAL,True)
table(s,['符号','含义','本器件中的注意点'],[
 ['Zt','传输线特性阻抗','由导体宽、缝隙和介质环境决定'],
 ['β、l','传播常数与长度','需要相应模式的有效波长'],
 ['ZL','被变换的负载','实际槽与相邻网络未必是纯电阻']],y=3.03,row_h=.73,col_widths=[1.9,3.3,6.88],size=18)
callout(s,'本文给出了电长度思路，但没有给出足以唯一重建模型的全部尺寸。')
# 11
image_slide('有限G–S–G轮廓：回流路径也是模型的一部分',REPO/'paper_layout_2021/layout_preview.png',[
 '展开焊盘便于连接；渐变区过渡到窄共面线。','地的边界会影响回流、寄生耦合和辐射。','旋转90°本身不是差异；矩形地替代真实轮廓才改变结构。'],source='图：本次重建的结构预览，图片估计尺寸；不是作者版图或求解结果。')
# 12
s=cards('宽窄交替的扼流网络为什么有效？',[
 ('阻抗阶跃','导体宽度和缝隙改变共面线特性阻抗；多个交替段共同构成频率选择网络。'),
 ('高频隔离','目标是在THz载频附近抑制能量沿读出端泄漏；正文描述交替λ/4段。'),
 ('读出需单独核对','允许直流通过不等于保证0–43 GHz无损通过。IF插损、反射及封装必须另行计算。')],notes='作者强调双功能射频线路。讲解时不要把低通滤波器简化成“所有低于130GHz都通过”。应建立实际终端下的传输响应，关注通带、阻带与寄生谐振。')
# 13
image_slide('齿形接触：让两侧响应不再相互抵消',A/'paper_fig1.png',[
 '一侧接触做成齿形，尖端附近局部场与吸收增强。','另一侧接触缺少同样增强，产生热电响应不对称。','几何不对称是零偏压设计关键，不等于必须外加偏置。'],source='来源：论文图1及正文；作者引用文献[19]支持结构化接触的先前工作。',notes='该页图像为论文原图，不是当前HFSS结果。正文认为另一接触的反向PTE因较弱吸收而减小，不能把“单侧齿形”仅解释成整流二极管。')
# 14
s=new('PTE模型：电磁场之后，还需要求温度与电势','多物理场',notes='以一般热电形式解释论文Methods的连续性方程。q_abs来自材料吸收，不是任意采用峰值|E|²；片层和体材料的定义需一致。参数S、热导、热容、能量弛豫时间与界面条件必须有来源。')
text(s,.75,2.05,11.7,.55,'V_PTE ∼ ∫ S(x) · ∇Te(x) · dl',27,TEAL,True)
bullets(s,['电磁：计算石墨烯中的吸收热源 q_abs。','热学：计算电子温度 Te；包含扩散、能量弛豫与边界散热。','电学：结合Seebeck系数与电流连续性，得到开路电压或短路电流。'],y=2.95,gap=.92,size=22)
callout(s,'一个集总端口可以提取馈点阻抗，但不能给出PTE响应率或零偏压光电压。')
# 15
s=new('材料堆叠：哪些是原文参数，哪些还缺依据？','材料与工艺')
table(s,['项目','原文信息','仿真处理与限制'],[
 ['石墨烯','电阻约1 kΩ','电磁负载近似≠完整复电导/PTE模型'],
 ['hBN封装','hBN/石墨烯/hBN','缺厚度和材料参数；不能凭图填值'],
 ['SiO₂','300 nm','新版显式加入；εr=3.9是初始假设'],
 ['硅','约100 Ω·cm','换算σ=1 S/m；εr=11.7为初值'],
 ['金属','Cr 3 nm / Au 80 nm','最终应检查薄膜损耗；结构核对版仍为PEC']],row_h=.69,col_widths=[1.85,3.55,6.68],size=17)
# 16
s=cards('硅与透镜：改变波长、耦合和损耗',[
 ('高折射率环境','硅改变传播常数和辐射分配。空气中的半波尺寸不能直接照搬到硅上天线。'),
 ('透镜耦合','论文把芯片置于硅透镜中心附近以耦合自由空间辐射；图2d计算是否含透镜需确认。'),
 ('导电衬底损耗','作者提醒6 μm缝隙中电场强；硅过于导电时，读出结构可能形成高频分流损耗。')],notes='不要把硅半空间、有限硅片和实体透镜当作完全等价。它们对驻波、辐射和端口阻抗的影响需要分别研究。')
# 17
s=new('旧模型与新版：改变的是电磁边界，而非外观装饰','模型修正')
pic(s,REPO/'paper_layout_2021/layout_preview.png',.65,1.95,11.9,2.95)
table(s,['旧矩形地基线','新版结构核对版'],[
 ['矩形地包围读出结构，焊盘轮廓简化','有限GSG地、展开焊盘与渐变、右端平台'],
 ['硅直接到金属面，氧化层省略','加入300 nm氧化层；内部界面不设辐射边界']],y=5.03,row_h=.49,col_widths=[6.04,6.04],size=15)
callout(s,'旧CSV不能代表新版；新版尚未在HFSS中求解，不能报告“已复现”。')
# 18
s=new('先冻结输入条件，才有资格比较论文曲线','复现输入')
table(s,['输入类别','必须记录','当前状态'],[
 ['几何','槽、槽距、馈点、扼流、焊盘、芯片','多数为图像估计'],
 ['材料','厚度、复介电常数、薄膜电导、石墨烯','正文只给出部分'],
 ['端口','位置、参考面、参考阻抗、另一端终端','原作者定义不充分'],
 ['环境','有限衬底、透镜、开放边界与封装','图2d具体计算条件待确认']],row_h=.84,col_widths=[1.65,6.15,4.28],size=18)
# 19
flow('HFSS工作流：每一步都应有可检查的输出',[
 ('几何','轮廓、层叠\n连通性、间隙'),('激励','端口与终端\n参考面、归一化'),('求解','网格、频率\n收敛与功率'),('比较','Z与S、场分布\n原图与误差')],
 '建议先建模检查，再启动求解；不要把验证通过当作计算精度保证。')
# 20
s=new('坐标与层叠：统一单位，避免“看起来对、实际错”','HFSS建模',notes='新版z坐标为金属/端口0，氧化层[-0.3,0]μm，硅顶面-0.3μm。空气1000μm、硅深1000μm是计算域假设，不是原文给出的硅厚度。矩形金属尺寸单位均为μm，80nm=0.08μm。')
for yy,hh,col,label0 in [(2.05,.9,'#E9F3F8','空气计算域：z > 0'),(3.0,.15,GOLD,'金属 / 端口：z = 0'),(3.23,.5,'#C9E4DD','SiO₂：−0.3 μm ≤ z ≤ 0'),(3.84,1.45,'#CBD5E0','硅：顶部 z = −0.3 μm')]:
 rect(s,.8,yy,7.2,hh,col)
 text(s,1,yy-.37 if hh<.4 else yy+.08,6.8,max(hh-.08,.35),label0,18,INK,True)
text(s,8.5,2.2,3.7,2.8,'示意图不按厚度比例绘制。\n\n300 nm = 0.3 μm\n80 nm = 0.08 μm\n100 Ω·cm = 1 Ω·m',21)
callout(s,'错误单位会同时改变谐振频率、端口长度与网格成本。')
# 21
s=new('几何参数表：给每个数字标明来源','HFSS建模')
table(s,['参数','当前初值','依据/注意'],[
 ['槽长 / 槽距','260 / 125 μm','显微图估计，非作者精确尺寸'],
 ['中心导体 / 端口间隙','12 / 6 μm','假设；端口间隙不是已知沟道长度'],
 ['宽窄段长度','每段120 μm，共10段','图像估计；电长度尚未核实'],
 ['宽段对地缝隙','6 μm','正文给出'],
 ['有限GSG焊盘与平台','按500 μm标尺重建','放大插图不可用作同尺度标尺']],row_h=.69,col_widths=[3.2,3.45,5.43],size=17)
# 22
s=cards('PEC是排查基线，不是最终薄膜材料模型',[
 ('PEC阶段','先检查拓扑和主导频率；计算中不计金属欧姆损耗，不能据此比较最终效率与峰深。'),
 ('薄膜阶段','原文Au仅80 nm。以体金σ≈4.1×10⁷ S/m估算，130 GHz趋肤深度约218 nm。'),
 ('建模选择','可研究适用的薄层/阻抗或有限导体模型；需核查两侧场、厚度、电导和Cr层影响。')],notes='趋肤深度δ=sqrt(2/(ωμσ))。218nm是采用体金电导、非磁性近似的理论估算，不是薄膜实测值。薄膜电导可能低于体材料。边界模型需核查适用条件，不能为贴合曲线任意调损耗。')
# 23
s=new('开放边界：让离开的波尽量不返回','HFSS边界',notes='旧两域模型外表面10面；新版空气/氧化层/硅三域外表面14面。面数是此特定几何的检查值，不是所有天线模型的规则。辐射边界不是完全无反射，尤其在介质侧、斜入射和近场区需要检验。')
bullets(s,['Radiation只赋给计算域外表面，不赋给内部空气/SiO₂、SiO₂/Si交界面。','在100 GHz，λ₀/4约750 μm；当前空气留白约900 μm，是起点而不是精度证明。','把侧向留白、空气高度、硅深度分别增大约50%，比较Zin和S11；必要时与合适PML对照。'],gap=1.17,size=21)
callout(s,'边界扩大后若曲线明显变化，先解决边界影响，再讨论“结构谐振”。')
# 24
s=new('端口有两种问题定义，不能混用','HFSS端口',notes='方案A的激励代表从石墨烯位置向天线送能以提取天线阻抗；源参考阻抗不等于计算石墨烯实际吸收。方案B是另一端激励下的带电阻网络，其S11不是方案A的同一量。原作者图注不足以唯一判定其全部设置。')
table(s,['方案','物理问题','应避免的错误'],[
 ['A：石墨烯位置集总端口','提取天线馈点Zin，再做1kΩ匹配分析','同一位置再并联1kΩ，导致重复加载'],
 ['B：保留1kΩ元件，另端激励','计算实际带负载网络的反射与传输','把另端口的S11当作石墨烯馈点S11']],row_h=1.12,col_widths=[3.2,4.4,4.48],size=18)
text(s,.8,5.66,11.5,.5,'当前基线采用A；作者实际端口与终端定义仍需核对。',22,TEAL,True)
# 25
s=new('集总端口检查：不仅是“能创建成功”','HFSS端口')
bullets(s,['端口片跨接两个不同导体；不能被赋成PEC，也不能落在短路区域。','积分线连接两侧导体。已修正为数值坐标：−3 μm → +3 μm，避免旧版解析变量为零长度。','方向决定电压正负参考；端口片尺寸和参考面也可能引入寄生，应做敏感性检查。','改动间隙或几何后，要同步更新积分线并重新验证。'],gap=.96,size=20)
callout(s,'脚本兼容性错误与物理模型错误，应分开排查。')
# 26
s=new('网格：优先细化电场变化快、功率交换强的区域','HFSS网格')
table(s,['区域','当前起始设置','需要验证的内容'],[
 ['馈点 / 6μm缝隙','端口最大长度2 μm','再细化至1 μm后阻抗是否稳定'],
 ['金属及槽边','最大长度12 μm','可对照8 μm，避免只追踪尖角峰值'],
 ['300nm氧化层','显式薄体，HFSS自动剖分','长宽厚尺度悬殊，网格成本与质量'],
 ['衬底 / 外部空间','波长与自适应网格控制','材料波长、远边界与局部误差']],row_h=.82,col_widths=[2.65,3.6,5.83],size=18)
# 27
s=new('Setup220：为什么先在频段上限自适应？','HFSS求解')
table(s,['设置','当前值','解释'],[
 ['自适应频率','220 GHz','高频波长较短，作为初始空间分辨率依据'],
 ['MaxDeltaS','0.02','相邻自适应轮次S变化判据，不是2%绝对误差'],
 ['轮数','至少3、最多15','达到最大轮数不等于收敛'],
 ['连续收敛轮数','2','要求变化连续满足判据']],row_h=.82,col_widths=[3,3.1,5.98],size=18)
callout(s,'220 GHz网格收敛不保证每个低频窄谐振均精确；关键频率仍需复核。')
# 28
s=new('用户基线：收敛通过，但不等于已复现论文','已有证据',source='来源：本会话用户上传的HFSS收敛截图；数值为截图读数，不是求解日志解析。')
table(s,['自适应轮次','Max |ΔS|','判断'],[
 ['1','N/A','没有上一轮可比较'],['2','约0.02137','略高于0.02'],
 ['3','约0.00731','满足判据'],['4','约0.00400','连续满足；界面显示CONVERGED']],row_h=.8,col_widths=[2.7,3.3,6.08],size=19)
callout(s,'这只证明旧模型的该自适应设置满足判据；新版还没有这样的求解证据。')
# 29
s=cards('两种扫频各做什么？',[
 ('插值扫频','100–220 GHz，输出241点。求解器选择部分频点建立近似；0.5 GHz是输出间隔，不代表逐点求解。'),
 ('离散核查','原设100、110…220 GHz共13点；之后用户加密155–170 GHz、间隔1 GHz。'),
 ('关键频率','旧模型162 GHz附近离散低谷约−16.2 dB，与插值趋势接近；不能推广为所有变体已核查。')],section='HFSS扫频',notes='CSV中260μm曲线最小值-15.83dB与此前截图约-16.2dB不完全相同，可能与求解/插值/模型状态有关。不能把两者说成同一原始运行精确复算。应保留各次工程与日志。')
# 30
s=new('怎样读S11与Zin：三条曲线要互相解释','结果解读')
text(s,.8,2.02,11.8,.6,'Γ = (Zin − Rref) / (Zin + Rref)，S11,dB = 20 log₁₀|Γ|',24,TEAL,True)
bullets(s,['同端口、同参考面、同终端且Rref为实数时，Zin = Rref(1+Γ)/(1−Γ)。','当前Rref=1000 Ω。较好的匹配通常要求实部靠近1000 Ω、虚部接近0。','−10 dB约为10%反射功率；1−|S11|²是净接受比例，不是石墨烯吸收率或辐射效率。'],y=3.05,gap=1.05,size=21)
callout(s,'HFSS阻抗通常以Ω显示；论文图2d用kΩ，比较前先统一单位。')
# 31 actual chart
s=new('槽长扫描：高频低谷移动，同时逐渐变浅','用户CSV分析',source='来源：用户 s11 SLOT.csv；旧矩形地模型；7组×241点，100–220 GHz。')
cd=XyChartData()
for length,rows in sorted(g.items()):
 ss=cd.add_series('%g μm'%length)
 for f0,v in rows:ss.add_data_point(f0,v)
ch=s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(.7),Inches(1.98),Inches(11.95),Inches(4.28),cd).chart
ch.has_legend=True;ch.legend.position=XL_LEGEND_POSITION.BOTTOM;ch.legend.font.size=Pt(11)
ch.value_axis.minimum_scale=-35;ch.value_axis.maximum_scale=0
ch.category_axis.minimum_scale=100;ch.category_axis.maximum_scale=220
ch.category_axis.tick_label_position=XL_TICK_LABEL_POSITION.LOW
for ax in [ch.value_axis,ch.category_axis]:ax.tick_labels.font.size=Pt(12);ax.tick_labels.font.name=FONT
palette=['204D78','7F9FBA','AD741B','CA9D49','6B733C','9B456E','D28EAA']
for se,c in zip(ch.series,palette):se.format.line.color.rgb=RGBColor.from_string(c);se.format.line.width=Pt(1.65)
ch.value_axis.major_gridlines.format.line.color.rgb=rgb(LINE)
ch.value_axis.major_gridlines.format.line.width=Pt(.5)
callout(s,'横轴：频率/GHz；纵轴：S11/dB。该图是用户旧模型结果，不是论文或新版结果。')
# 32 table actual
s=new('七组数据：分别记录目标分支与130 GHz处表现','用户CSV分析',source='来源：s11 SLOT.csv；高频分支按140–200 GHz局部低谷跟踪，并非已证明的模式身份。')
rows=[[str(int(l)), '%.1f'%fh,'%.2f'%sh,'%.2f'%s130] for l,fh,sh,fl,sl,s130 in summary]
table(s,['槽长/μm','高频分支低谷/GHz','该低谷S11/dB','130 GHz处S11/dB'],rows,row_h=.48,col_widths=[2,3.4,3.3,3.38],size=16)
callout(s,'220→340 μm：低谷170→144.5 GHz，但其S11由−31.17→−6.69 dB。')
# 33 two branches
s=new('340 μm的“全局最小值”已经换了分支','用户CSV分析',source='来源：s11 SLOT.csv；各频率为0.5 GHz导出网格上的局部最低采样点。',notes='在340μm时，129.5GHz处-6.7857dB比144.5GHz处-6.6933dB略深。若只记录argmin会产生从149到129.5GHz的假跳变，不代表同一模式连续移动到130GHz。')
cd=XyChartData()
for label0,index in [('原高频低谷分支',1),('约130 GHz低谷分支',3)]:
 ss=cd.add_series(label0)
 for row in summary:ss.add_data_point(row[0],row[index])
ch=s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, Inches(.65),Inches(1.98),Inches(7.65),Inches(4.25),cd).chart
ch.has_legend=True;ch.legend.position=XL_LEGEND_POSITION.BOTTOM;ch.legend.font.size=Pt(11)
ch.category_axis.minimum_scale=220;ch.category_axis.maximum_scale=340
ch.value_axis.minimum_scale=120;ch.value_axis.maximum_scale=180
ch.value_axis.major_gridlines.format.line.color.rgb=rgb(LINE)
ch.value_axis.major_gridlines.format.line.width=Pt(.5)
for ax in [ch.value_axis,ch.category_axis]:ax.tick_labels.font.size=Pt(12)
for se,c in zip(ch.series,['204D78','9B456E']):se.format.line.color.rgb=RGBColor.from_string(c);se.marker.size=7
bullets(s,['340 μm：129.5 GHz约−6.79 dB。','同组另一低谷：144.5 GHz约−6.69 dB。','分支身份需要场图与电流分布确认。'],x=8.6,y=2.25,w=4.0,gap=1.12,size=19)
callout(s,'图横轴：槽长/μm；纵轴：低谷频率/GHz。不能把所有全局最小值连成同一个模式。')
# 34
s=new('怎样判断低谷来自双槽，还是来自读出网络？','诊断方法')
table(s,['检查量','双槽主导的可能表现','读出/扼流主导的可能表现'],[
 ['场与表面电流','槽边、馈点附近变化显著','宽窄段或焊盘处强驻波'],
 ['参数敏感性','槽长、槽距变化影响较大','段长、焊盘或IF终端变化敏感'],
 ['功率去向','辐射/敏感区耦合占比变化','沿线泄漏或衬底损耗显著']],row_h=.94,col_widths=[2.2,4.85,5.03],size=18)
callout(s,'这些是诊断线索，不是仅凭一张S11图即可确定的模式结论。')
# 35
s=new('下一轮扫参：先核对结构，再控制变量','参数研究',notes='旧CSV仅用于说明旧模型对槽长敏感，不能把其最优槽长直接迁移到新有限地模型。每次改变终端或材料需要重新建立基准。重点同时比较ReZ、ImZ、S11和模式分布。')
table(s,['阶段','变量/条件','目的'],[
 ['结构基准','固定原图重建轮廓与材料假设','建立新版独立基线，不沿用旧模型结果'],
 ['几何敏感性','槽长、槽距、馈点距离分别扫描','跟踪分支与阻抗变化'],
 ['网络影响','扼流段长度、IF开路/实际终端','识别网络谐振和负载牵引'],
 ['模型稳健性','网格、边界、损耗、衬底条件','区分物理变化与数值/假设影响']],row_h=.82,col_widths=[2.2,4.6,5.28],size=18)
# 36
flow('完整探测性能：把电磁吸收传给热电模型',[
 ('电磁','E场与复电导\n→实际吸收热源'),('热学','扩散与能量弛豫\n→电子温度'),('热电','接触与Seebeck分布\n→光电压/电流'),('校准','入射功率定义\n→响应率与NEP')],
 'Pinc、透镜收集功率、天线接受功率与石墨烯吸收功率不能混用。',section='多物理场扩展',notes='论文响应率采用入射功率归一化。若模拟只用石墨烯吸收功率作分母，得到的是另一种内响应率，不能与图2b直接比较。NEP还需要噪声模型与测量带宽信息。')
# 37
s=cards('“超过43 GHz”还需要一套读出与瞬态验证',[
 ('局部响应速度','电子冷却与分布式RC。靠近接触产生响应有望缩短传播延迟；本文引用文献[25]。'),
 ('外部网络','CPW、焊线、PCB、放大器和连接器可能限制最终电学带宽。器件快不等于整机快。'),
 ('差频实验','扫描IF观察输出响应。实际频响与仪器链路需校准，不能用载频S11代替。')],section='高速读出',source='依据：论文图3、第5页讨论、Methods；文献[25]为作者引用的RC依据。')
# 38
s=new('什么证据才支持“复现了图2d”？','验收标准')
table(s,['层级','需要的证据','不能替代它的现象'],[
 ['几何与条件一致','尺寸/材料/端口/终端/环境可追溯','照片轮廓看起来相近'],
 ['数值稳定','收敛、网格/边界复核、离散点与功率检查','一次Validate通过或求解正常结束'],
 ['频谱一致','同端口的ReZ、ImZ与S11整体对比','只让一个低谷落在130 GHz'],
 ['独立可复算','保存工程、参数、原始CSV与日志','只保留截图或最终拟合参数']],row_h=.82,col_widths=[2.3,5.1,4.68],size=18)
# 39
s=new('当前项目状态与优先工作','执行清单')
table(s,['状态','已经具备','下一步'],[
 ['旧模型：已求解','用户收敛截图与7组槽长CSV','保留为基线，停止当作原文版图'],
 ['新版：结构核对','有限GSG轮廓、300nm氧化层、二维连通检查','先在2021 R1建模与求解验证'],
 ['物理细化：未完成','原文部分材料与工艺参数','薄膜损耗、真实终端、衬底/透镜敏感性'],
 ['精确复现：仍受阻','正文与显微图','需要作者尺寸/CAD、端口和终端等条件']],row_h=.82,col_widths=[2.8,4.7,4.58],size=17)
# 40
s=new('参考依据：原文、其引用与用户数据分别使用','来源与参考',notes='本PPT直接阅读来源是用户上传主文PDF与用户CSV。下列[19][20][21][25]的书目信息和引用目的由主文确认，未声称已阅读其全文。原文数据可按Data Availability从补充信息或通讯作者获取，本任务未取得CAD、补充尺寸或端口定义。')
refs=[
 '主文：Titova等，Graphene Zero-Bias Sub-Terahertz Turnkey Detector with Above 43 GHz Bandwidth，arXiv:2603.03554v3。',
 '[19] Semkin等，Advanced Optical Materials 13 (2025)，DOI: 10.1002/adom.202403189；结构化接触与零偏压依据。',
 '[20] Filipovic, Gearhart, Rebeiz，IEEE T-MTT 41, 1738–1749 (1993)；双槽透镜天线依据。',
 '[21] Zhang等，IEEE T-THz 1, 378–382 (2011)；探测器与双槽天线耦合背景。',
 '[25] Safonov, Svintsov，Response times of two-dimensional photodetectors limited by intrinsic resistance and capacitance (2025)。',
 '用户数据：s11 SLOT.csv（7组槽长）；本会话HFSS截图。理论式为常规电磁、传输线和热电关系。']
bullets(s,refs,y=1.98,gap=.72,size=16)
# 41 appendix
s=new('附录：对照论文图2，不把不同指标混在一起','论文对照')
pic(s,A/'paper_fig2.png',.6,1.9,7.6,4.95)
bullets(s,['图2b：响应率与噪声等效功率。','图2c：实测载频响应，包含光学耦合与器件。','图2d：计算的输入阻抗和S参数；不是完整响应率模型。'],x=8.5,y=2.1,w=4.0,gap=1.35,size=19)
# 42
s=new('附录：保存与导出的最小记录包','操作记录')
table(s,['文件/截图','保存内容','用途'],[
 ['.aedt及结果目录','模型、材料、边界、各次变体','重建和复算'],
 ['参数表与版本说明','所有尺寸、估计来源、改变项','避免混淆旧/新版与不同终端'],
 ['复数S与Z的CSV','频率、变体、实部/虚部、单位','独立校验归一化与阻抗关系'],
 ['日志与收敛/场图','每组状态、网格、功率、关键场分布','证明计算完成、稳定且物理可解释']],row_h=.82,col_widths=[2.7,4.7,4.68],size=18)
callout(s,'建议使用英文工程名和独立运行目录；先归档基线，再修改几何或材料。')

out=R/'Graphene_THz_HFSS_Simulation_Guide.pptx'
prs.core_properties.title='石墨烯零偏压太赫兹探测器：结构设计与仿真思路'
prs.core_properties.subject='HFSS 2021 R1；论文解读与复现验证'
prs.core_properties.author='Codex'
prs.core_properties.keywords='graphene,THz,HFSS,PTE,double-slot,simulation'
prs.save(out)
(R/'slide_manifest.json').write_text(json.dumps(slides_meta,ensure_ascii=False,indent=2),encoding='utf-8')
(R/'layout_checks.json').write_text(json.dumps({'slides':len(prs.slides),'text_boxes':len(checks),'text_fit':'passed explicit font-metric bounds','checks':checks},ensure_ascii=False,indent=2),encoding='utf-8')
(R/'speaker_notes.md').write_text('\n\n'.join('## %02d %s\n\n%s\n\n%s'%(v['page'],v['title'],v.get('notes','') or '按页面结构讲解，强调事实、假设与证据的区别。',v['source']) for v in slides_meta),encoding='utf-8')
print('Created',out,'slides',len(prs.slides),'bytes',out.stat().st_size)
